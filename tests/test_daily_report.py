import io
import docx
import pytest
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient

import auth
import db
import main

client = TestClient(main.app)


def test_daily_report_requires_auth():
    response = client.get("/reports/daily?date=2026-09-08")
    assert response.status_code == 401


def test_daily_report_requires_admin_role():
    with patch.object(auth, "verify_jwt", return_value={"sub": "user-member"}), \
         patch.object(db, "get_profile_by_id", return_value={"id": "user-member", "client_id": "client-1", "role": "member"}):
        headers = {"Authorization": "Bearer token"}
        response = client.get("/reports/daily?date=2026-09-08", headers=headers)
        assert response.status_code == 403
        assert "Admin role required" in response.json()["detail"]


def test_daily_report_rejects_invalid_date_format():
    with patch.object(auth, "verify_jwt", return_value={"sub": "user-admin"}), \
         patch.object(db, "get_profile_by_id", return_value={"id": "user-admin", "client_id": "client-1", "role": "admin"}):
        headers = {"Authorization": "Bearer token"}
        response = client.get("/reports/daily?date=invalid-date", headers=headers)
        assert response.status_code == 400
        assert "Invalid date format" in response.json()["detail"]


def test_daily_report_returns_no_data_when_zero_rfqs():
    with patch.object(auth, "verify_jwt", return_value={"sub": "user-admin"}), \
         patch.object(db, "get_profile_by_id", return_value={"id": "user-admin", "client_id": "client-1", "role": "admin"}), \
         patch.object(db, "get_rfqs_by_date", return_value=[]):
        headers = {"Authorization": "Bearer token"}
        response = client.get("/reports/daily?date=2026-09-08", headers=headers)

        assert response.status_code == 200
        data = response.json()
        assert data["status"] == "no_data"
        assert data["message"] == "No RFQs were created on this date."
        assert "application/json" in response.headers.get("content-type", "")


def test_daily_report_generates_docx_with_mix_of_responded_and_unresponded_rfqs():
    mock_rfqs = [
        {
            "id": "rfq-1",
            "client_id": "client-1",
            "product_name": "Aluminium Ladder",
            "category": "Hardware",
            "quantity": 10,
            "specs": "4x5 Heavy Duty",
            "deadline_hours": 24,
            "created_at": "2026-09-08T10:00:00Z",
        },
        {
            "id": "rfq-2",
            "client_id": "client-1",
            "product_name": "Power Drill 18V",
            "category": "Power Tools",
            "quantity": 5,
            "specs": "Brushless Motor",
            "deadline_hours": 12,
            "created_at": "2026-09-08T14:30:00Z",
        },
    ]

    mock_quotes_rfq1 = [
        {
            "id": "q-1",
            "supplier_id": "sup-1",
            "price": 120.0,
            "delivery_time": "2 days",
            "quality_notes": "1-year warranty",
            "suppliers": {"name": "Ace Hardware"},
        },
        {
            "id": "q-2",
            "supplier_id": "sup-2",
            "price": 110.0,
            "delivery_time": "1 day",
            "quality_notes": "Standard",
            "suppliers": {"name": "BuildPro Supplies"},
        },
    ]

    mock_ranking_rfq1 = {
        "best_supplier_id": "sup-2",
        "reasoning": "BuildPro Supplies offers the lowest price (AED 110.0) with fastest delivery (1 day).",
        "ranking": [
            {"supplier_id": "sup-2", "rank": 1, "summary": "Best price AED 110.0 and 1-day delivery."},
            {"supplier_id": "sup-1", "rank": 2, "summary": "AED 120.0 with 2-day delivery."},
        ],
    }

    def fake_get_quotes(rfq_id):
        if rfq_id == "rfq-1":
            return mock_quotes_rfq1
        return []

    def fake_get_ranking(rfq_id):
        if rfq_id == "rfq-1":
            return {"ranking_json": mock_ranking_rfq1}
        return None

    with patch.object(auth, "verify_jwt", return_value={"sub": "user-admin"}), \
         patch.object(db, "get_profile_by_id", return_value={"id": "user-admin", "client_id": "client-1", "role": "admin"}), \
         patch.object(db, "get_rfqs_by_date", return_value=mock_rfqs), \
         patch.object(db, "get_quotes_for_rfq", side_effect=fake_get_quotes), \
         patch.object(db, "get_ranking_for_rfq", side_effect=fake_get_ranking):

        headers = {"Authorization": "Bearer token"}
        response = client.get("/reports/daily?date=2026-09-08", headers=headers)

        assert response.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in response.headers.get("content-type", "")
        assert 'Daily_Procurement_Report_2026-09-08.docx' in response.headers.get("content-disposition", "")

        # Parse the docx to verify contents
        doc = docx.Document(io.BytesIO(response.content))
        doc_text = "\n".join(p.text for p in doc.paragraphs)

        assert "Daily Procurement Report — 2026-09-08" in doc_text
        assert "RFQ: Aluminium Ladder" in doc_text
        assert "RFQ: Power Drill 18V" in doc_text
        assert "No one responded to this RFQ." in doc_text

        # Verify tables
        assert len(doc.tables) == 1
        table = doc.tables[0]
        rows = [[c.text for c in row.cells] for row in table.rows]
        assert rows[0] == ["Rank", "Supplier", "Price", "Delivery Time", "Quality / Warranty Notes"]
        # Rank #1 should be BuildPro Supplies (AED 110.0)
        assert rows[1][0] == "#1"
        assert rows[1][1] == "BuildPro Supplies"
        assert rows[1][2] == "AED 110.0"
        # Rank #2 should be Ace Hardware (AED 120.0)
        assert rows[2][0] == "#2"
        assert rows[2][1] == "Ace Hardware"
        assert rows[2][2] == "AED 120.0"
